"""
GMAT Club -- CR x Source: OG  Full Scraper  (with Login)
=========================================================
Scrapes ALL questions tagged with any "Source: OG *" tag
inside the Critical Reasoning (CR) category.

Phase 1 -- Link collection (same approach as PS scraper):
  Paginates through tag search results, collects question links.

Phase 2 -- Content extraction:
  Visits each question page, extracts passage, question stem,
  options A-E, correct answer, source, and tags.

Login Strategy
--------------
1. Browser opens to GMAT Club login page
2. YOU log in manually (handles Cloudflare + 2FA naturally)
3. Script detects when login is complete and starts scraping automatically
4. Session is saved to disk -- next run skips login entirely

What it does
------------
1. Reads gmatclub_tags.xlsx to get the CR Source:OG tag IDs
2. For each tag ID, paginates through ALL search result pages (start=0,50,100...)
3. From every question card captures: Title, Link, Category, Tags, Tag IDs
4. Deduplicates globally -- a question under multiple OG tags is saved only once
5. Visits each question page and extracts:
      Question Stem, Question + Text, A, B, C, D, E, Answer, Source
6. Saves a master Excel:
      - Sheet "All Questions"  -- full deduplicated list with all content
      - Sheet per OG tag label -- e.g. "OG 2025-2026", "OG 2022" ...
      - Sheet "Summary"        -- count per tag + grand total

Output
------
  CR_OG_questions.xlsx

Columns
-------
  No | Title | Link | Category | Tags | Tag IDs | Source |
  Question Stem | Question + Text | A | B | C | D | E | Answer

Usage
-----
  python scrape_cr_og.py                              # normal run (login if needed)
  python scrape_cr_og.py --fresh-login                # force re-login even if session exists
"""

import argparse
import asyncio
import json
import re
import sys
from pathlib import Path
from urllib.parse import urlencode

try:
    from bs4 import BeautifulSoup
    import pandas as pd
    from openpyxl.styles import Font, PatternFill, Alignment
    from playwright.async_api import async_playwright
except ImportError:
    sys.exit(
        "Missing dependencies. Run:\n"
        "  pip install beautifulsoup4 pandas openpyxl playwright\n"
        "  playwright install chromium"
    )

# ── Config ────────────────────────────────────────────────────────────────────

TAGS_EXCEL = "gmatclub_tags.xlsx"
OUTPUT_EXCEL = "CR_OG_questions.xlsx"
SESSION_FILE = "gmatclub_session.json"  # saved cookies -- reused across runs
LOGIN_URL = "https://gmatclub.com/forum/ucp.php?mode=login"
FORUM_HOME = "https://gmatclub.com/forum/"
BASE_SEARCH_URL = "https://gmatclub.com/forum/search.php"
PAGE_SIZE = 50
NO_RESULTS_MSG = "No suitable matches were found."
CF_WAIT_MS = 4000
CF_RETRIES = 20
PAGE_TIMEOUT = 60000
CRAWL_DELAY = 2  # seconds between search-result page requests
CONTENT_DELAY = 1.5  # seconds between individual question page fetches

COL_ORDER = [
    "No",
    "Title",
    "Link",
    "Category",
    "Tags",
    "Tag IDs",
    "Source",
    "Question Stem",
    "Question + Text",
    "A",
    "B",
    "C",
    "D",
    "E",
    "Answer",
]

COL_WIDTHS = {
    "No": 6,
    "Title": 45,
    "Link": 30,
    "Category": 22,
    "Tags": 55,
    "Tag IDs": 25,
    "Source": 22,
    "Question Stem": 60,
    "Question + Text": 80,
    "A": 55,
    "B": 55,
    "C": 55,
    "D": 55,
    "E": 55,
    "Answer": 10,
}


# ── Load Target Tags ──────────────────────────────────────────────────────────


def load_cr_og_tags(tags_excel: str) -> list:
    df = pd.read_excel(tags_excel, sheet_name="All Tags")
    mask = (df["Category"] == "Critical Reasoning (CR)") & (
        df["Tag Label"].str.startswith("Source: OG")
    )
    rows = df[mask][["Tag ID", "Tag Label"]].reset_index(drop=True)
    tags = [
        {"tag_id": int(r["Tag ID"]), "tag_label": str(r["Tag Label"])}
        for _, r in rows.iterrows()
    ]
    print(f"\n🏷   Found {len(tags)} CR Source:OG tags to scrape:")
    for t in tags:
        print(f"       {t['tag_id']:>6}  {t['tag_label']}")
    return tags


# ── URL Builder ───────────────────────────────────────────────────────────────


def build_url(tag_id: int, start: int = 0) -> str:
    params = [
        ("selected_search_tags[]", tag_id),
        ("search_tags", "exact"),
        ("submit", "Search"),
    ]
    if start:
        params.append(("start", start))
    return BASE_SEARCH_URL + "?" + urlencode(params)


# ── Search Page Parser ────────────────────────────────────────────────────────


def parse_search_page(html: str):
    """
    Returns (questions_list, is_done).
    is_done=True means no more pages for this tag.
    """
    if NO_RESULTS_MSG in html:
        return [], True

    soup = BeautifulSoup(html, "html.parser")
    cards = soup.select("div.topicsName")
    if not cards:
        return [], False

    questions = []
    for card in cards:
        link_tag = card.select_one("a.topic-link")
        if not link_tag:
            continue

        href = link_tag.get("href", "").strip()
        title = link_tag.get("title", "").strip()
        if not title:
            span = link_tag.select_one("span.topicTitle")
            title = span.get_text(strip=True) if span else ""
        if href.startswith("/"):
            href = "https://gmatclub.com" + href

        cat_tag = card.select_one("p > a")
        category = cat_tag.get_text(strip=True) if cat_tag else ""

        tags_div = card.select_one("div.topic-tags")
        tag_labels = []
        tag_ids = []
        if tags_div:
            for a in tags_div.select("a[href*='tag_id=']"):
                lbl = a.get_text(strip=True)
                m = re.search(r"tag_id=(\d+)", a.get("href", ""))
                if m:
                    tag_labels.append(lbl)
                    tag_ids.append(m.group(1))

        questions.append(
            {
                "Title": title,
                "Link": href,
                "Category": category,
                "Tags": " | ".join(tag_labels),
                "Tag IDs": " | ".join(tag_ids),
            }
        )

    return questions, False


# ── Question Page Parser ──────────────────────────────────────────────────────

# Matches an option line: letter A-E, followed by zero-width space / nbsp / regular space
_OPT_RE = re.compile(r"^([A-E])[​  \t]+(.+)$")


def _clean_text(text: str) -> str:
    """Strip zero-width characters and collapse excessive whitespace."""
    text = text.replace("​", "").replace("‏", "").replace("­", "")
    text = re.sub(r"[ \t]+", " ", text)  # collapse horizontal whitespace
    text = re.sub(r"\n{3,}", "\n\n", text)  # at most double newlines
    return text.strip()


def _empty_content(source: str = "", tags: str = "", tag_ids: str = "") -> dict:
    return {
        "Question Stem": "",
        "Question + Text": "",
        "A": "",
        "B": "",
        "C": "",
        "D": "",
        "E": "",
        "Answer": "",
        "Source": source,
        "Tags": tags,
        "Tag IDs": tag_ids,
    }


def parse_question_page(html: str) -> dict:
    """
    Extract CR question content from a GMAT Club question page.

    Returns a dict with keys:
      Question Stem, Question + Text, A, B, C, D, E, Answer, Source, Tags, Tag IDs
    """
    soup = BeautifulSoup(html, "html.parser")

    # ── Tags (from tags-wrapper at bottom of page) ────────────────────────────
    taglist_div = soup.select_one("div#taglist")
    tag_labels, tag_ids, source = [], [], ""
    if taglist_div:
        for a in taglist_div.select("a.tag_css_link"):
            lbl = a.get_text(strip=True)
            m = re.search(r"tag_id=(\d+)", a.get("href", ""))
            if lbl and m:
                tag_labels.append(lbl)
                tag_ids.append(m.group(1))
            if lbl.startswith("Source:") and not source:
                source = lbl

    tags_str = " | ".join(tag_labels)
    tag_ids_str = " | ".join(tag_ids)

    # ── First post content ────────────────────────────────────────────────────
    item_text = soup.select_one("div.item.text")
    if not item_text:
        return _empty_content(source, tags_str, tag_ids_str)

    # Extract answer from spoiler before we decompose anything
    spoiler = item_text.select_one("div[id^='spoiler_']")
    raw_answer = spoiler.get_text(strip=True) if spoiler else ""
    # Grab only the first capital letter (handles "B " or "B or C" etc.)
    answer_match = re.search(r"[A-E]", raw_answer)
    answer = answer_match.group(0) if answer_match else raw_answer.strip()

    # Remove non-question elements before text extraction
    for sel in (
        "div.twoRowsBlock",
        "div.post_signature",
        "div.answer-block",
        "script",
        "style",
    ):
        for el in item_text.select(sel):
            el.decompose()

    # Convert <br> tags to newlines so line-based parsing works
    for br in item_text.find_all("br"):
        br.replace_with("\n")

    raw_text = item_text.get_text(separator="")
    text = _clean_text(raw_text)

    # ── Identify option lines ─────────────────────────────────────────────────
    lines = [l.strip() for l in text.split("\n")]

    options = {}  # {"A": "text", "B": "text", ...}
    opt_start_i = None

    for i, line in enumerate(lines):
        m = _OPT_RE.match(line)
        if m:
            letter = m.group(1)
            opt_text = m.group(2).strip()
            if letter not in options:  # first match per letter wins
                options[letter] = opt_text
                if opt_start_i is None:
                    opt_start_i = i

    # ── Split passage and question stem ───────────────────────────────────────
    pre_lines = lines[:opt_start_i] if opt_start_i is not None else lines
    pre_text = "\n".join(pre_lines)

    # Paragraphs = sections separated by one or more blank lines
    paragraphs = [p.strip() for p in re.split(r"\n{2,}", pre_text) if p.strip()]

    if len(paragraphs) >= 2:
        question_stem = paragraphs[-1]
        passage = "\n\n".join(paragraphs[:-1])
    elif paragraphs:
        question_stem = paragraphs[0]
        passage = ""
    else:
        question_stem = ""
        passage = ""

    # ── Build "Question + Text" (full passage + stem + all options) ───────────
    option_block = "\n".join(
        f"{letter} {options[letter]}" for letter in "ABCDE" if options.get(letter)
    )
    full_text_parts = [p for p in [passage, question_stem, option_block] if p]
    full_text = "\n\n".join(full_text_parts)

    return {
        "Question Stem": question_stem,
        "Question + Text": full_text,
        "A": options.get("A", ""),
        "B": options.get("B", ""),
        "C": options.get("C", ""),
        "D": options.get("D", ""),
        "E": options.get("E", ""),
        "Answer": answer,
        "Source": source,
        "Tags": tags_str,
        "Tag IDs": tag_ids_str,
    }


# ── Cloudflare Wait ───────────────────────────────────────────────────────────


async def wait_for_cloudflare(page) -> bool:
    for attempt in range(CF_RETRIES):
        content = await page.content()
        if (
            "Just a moment" not in content
            and "security verification" not in content.lower()
        ):
            return True
        print(f"    ⏳ Cloudflare... ({attempt + 1}/{CF_RETRIES})")
        await page.wait_for_timeout(CF_WAIT_MS)
    return False


# ── Login Handler ─────────────────────────────────────────────────────────────


def is_logged_in(html: str) -> bool:
    """Detect if the current page shows a logged-in user."""
    indicators = [
        "logout",
        "ucp.php?mode=logout",
        'class="icon-svg-logout"',
        "My Profile",
    ]
    html_lower = html.lower()
    return any(ind.lower() in html_lower for ind in indicators)


async def login(page, session_file: str, fresh_login: bool = False) -> bool:
    """
    Attempt to restore session from disk first.
    If that fails (or fresh_login=True), open login page for manual login.
    Returns True if logged in successfully.
    """
    session_path = Path(session_file)

    # ── Try restoring saved session ──────────────────────────────────────────
    if not fresh_login and session_path.exists():
        print("\n🔑  Found saved session -- attempting to restore...")
        try:
            cookies = json.loads(session_path.read_text())
            await page.context.add_cookies(cookies)
            await page.goto(FORUM_HOME, timeout=PAGE_TIMEOUT)
            await wait_for_cloudflare(page)
            html = await page.content()
            if is_logged_in(html):
                print("    ✅  Session restored -- already logged in!")
                return True
            else:
                print("    ⚠️  Saved session expired -- need to log in again.")
        except Exception as e:
            print(f"    ⚠️  Could not restore session: {e}")

    # ── Manual login ─────────────────────────────────────────────────────────
    print("\n🔐  Opening GMAT Club login page...")
    print("    👉  Please log in manually in the browser window.")
    print("    👉  The script will continue automatically once you are logged in.")
    print("    (You have up to 3 minutes)\n")

    await page.goto(LOGIN_URL, timeout=PAGE_TIMEOUT)
    await wait_for_cloudflare(page)

    for i in range(36):  # 36 x 5s = 180s = 3 minutes
        html = await page.content()
        if is_logged_in(html):
            print("\n    ✅  Login detected! Starting scrape...\n")
            cookies = await page.context.cookies()
            session_path.write_text(json.dumps(cookies, indent=2))
            print(f"    💾  Session saved to {session_file} (reused on next run)")
            return True
        remaining = (36 - i) * 5
        print(f"    ⏳  Waiting for login... ({remaining}s remaining)", end="\r")
        await page.wait_for_timeout(5000)

    print("\n    ❌  Login timeout. Please run the script again.")
    return False


# ── Phase 1: Collect Links per Tag ────────────────────────────────────────────


async def collect_links_for_tag(page, tag: dict, seen_links: set) -> list:
    tag_id = tag["tag_id"]
    tag_label = tag["tag_label"]
    collected = []
    start = 0
    page_num = 1

    print(f"\n  ── {tag_label}  (id={tag_id}) ──────────────────────────")

    while True:
        url = build_url(tag_id, start)
        print(f"    Page {page_num}  start={start}")

        try:
            await page.goto(url, timeout=PAGE_TIMEOUT)
        except Exception as e:
            print(f"    ✗ Navigation failed: {e}")
            break

        if not await wait_for_cloudflare(page):
            print("    ⚠️  Cloudflare did not clear -- skipping this tag.")
            break

        # Safety check: did we get logged out mid-scrape?
        html = await page.content()
        if not is_logged_in(html):
            print("    ⚠️  Session expired mid-scrape! Please re-run the script.")
            break

        rows, done = parse_search_page(html)

        if done:
            print(f"    ✅  No more results for this tag.")
            break

        if not rows:
            print(f"    ⚠️  Zero cards found -- stopping this tag.")
            break

        new_rows = [r for r in rows if r["Link"] not in seen_links]
        seen_links.update(r["Link"] for r in new_rows)
        collected.extend(new_rows)

        dupes = len(rows) - len(new_rows)
        print(
            f"    ✓  {len(new_rows):>3} new  |  {dupes:>3} dupes  |  tag total: {len(collected)}"
        )

        if len(rows) < PAGE_SIZE:
            print(f"    ✅  Last page.")
            break

        start += PAGE_SIZE
        page_num += 1
        await asyncio.sleep(CRAWL_DELAY)

    return collected


# ── Phase 2: Fetch Question Content ──────────────────────────────────────────


async def fetch_all_content(page, rows: list) -> None:
    """
    Visit each question URL and enrich the row dict with content fields.
    Mutates rows in-place.
    """
    total = len(rows)
    print(f"\n{'='*60}")
    print(f"  Phase 2: Fetching content for {total} questions...")
    print(f"{'='*60}")

    for i, row in enumerate(rows, start=1):
        url = row["Link"]
        print(f"  [{i:>4}/{total}]  {url}", end="  ")

        try:
            await page.goto(url, timeout=PAGE_TIMEOUT)

            if not await wait_for_cloudflare(page):
                print("⚠️  CF timeout -- skipping")
                _fill_empty_content(row)
                continue

            html = await page.content()
            content = parse_question_page(html)

            # Tags from the question page are more complete than search-result tags
            if content.get("Tags"):
                row["Tags"] = content["Tags"]
                row["Tag IDs"] = content["Tag IDs"]

            for field in (
                "Question Stem",
                "Question + Text",
                "A",
                "B",
                "C",
                "D",
                "E",
                "Answer",
                "Source",
            ):
                row[field] = content.get(field, "")

            has_all_opts = all(row.get(l) for l in "ABCDE")
            print(
                f"✓  Answer={row['Answer'] or '?'}  "
                f"Options={'ABCDE' if has_all_opts else 'partial'}"
            )

        except Exception as e:
            print(f"✗  {e}")
            _fill_empty_content(row)

        await asyncio.sleep(CONTENT_DELAY)


def _fill_empty_content(row: dict) -> None:
    """Ensure all content fields exist in a row (used on fetch failure)."""
    for field in (
        "Question Stem",
        "Question + Text",
        "A",
        "B",
        "C",
        "D",
        "E",
        "Answer",
        "Source",
    ):
        row.setdefault(field, "")


# ── Master Orchestrator ───────────────────────────────────────────────────────


async def run_scraper(tags: list, output: str, session_file: str, fresh_login: bool):
    all_rows = []
    seen_links = set()
    tag_buckets = {}

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=False,
            args=[
                "--no-sandbox",
                "--disable-dev-shm-usage",
                "--disable-blink-features=AutomationControlled",
            ],
        )
        context = await browser.new_context()
        page = await context.new_page()

        # Stealth
        await page.add_init_script(
            "Object.defineProperty(navigator,'webdriver',{get:()=>undefined})"
        )

        # Login
        logged_in = await login(page, session_file, fresh_login)
        if not logged_in:
            await browser.close()
            return

        # ── Phase 1: collect links ────────────────────────────────────────────
        print(f"\n{'='*60}")
        print(f"  Phase 1: Collecting question links...")
        print(f"{'='*60}")

        for tag in tags:
            rows = await collect_links_for_tag(page, tag, seen_links)
            tag_buckets[tag["tag_label"]] = rows
            all_rows.extend(rows)
            print(f"    📦  {tag['tag_label']}: {len(rows)} unique questions")

        print(f"\n  Total unique questions found: {len(all_rows)}")

        # ── Phase 2: fetch content ────────────────────────────────────────────
        await fetch_all_content(page, all_rows)

        # tag_buckets hold references to the same dicts that all_rows mutated,
        # so per-tag sheets will automatically include the content fields.

        await browser.close()

    print(f"\n{'='*60}")
    print(f"  Total unique questions scraped: {len(all_rows)}")
    print(f"{'='*60}")
    save_excel(all_rows, tag_buckets, output)


# ── Excel Saver ───────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
ALT_FILL = PatternFill("solid", fgColor="EBF3FB")


def _style_sheet(ws, df: pd.DataFrame):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    link_col = None
    for idx, col_cell in enumerate(ws[1], start=1):
        if col_cell.value == "Link":
            link_col = idx
            break

    for row_idx in range(2, ws.max_row + 1):
        if row_idx % 2 == 0:
            for cell in ws[row_idx]:
                cell.fill = ALT_FILL
        if link_col:
            cell = ws.cell(row=row_idx, column=link_col)
            if cell.value and str(cell.value).startswith("http"):
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")

    for col_cells in ws.columns:
        header = col_cells[0].value
        width = COL_WIDTHS.get(header, 18)
        ws.column_dimensions[col_cells[0].column_letter].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def _build_df(rows: list, start_no: int = 1) -> pd.DataFrame:
    df = pd.DataFrame(rows)
    # Ensure every expected column exists
    for col in COL_ORDER[1:]:  # skip "No" -- added below
        if col not in df.columns:
            df[col] = ""
    df.insert(0, "No", range(start_no, start_no + len(df)))
    return df[COL_ORDER]


def _sanitize_sheet_name(raw_name: str, existing_names: set[str]) -> str:
    sanitized = re.sub(r"[:\\/?*\[\]]+", " ", raw_name.replace("Source: ", "")).strip()
    if not sanitized:
        sanitized = "Tag"
    sanitized = sanitized[:31]

    base_name = sanitized
    idx = 1
    while sanitized in existing_names:
        suffix = f"_{idx}"
        truncated = base_name[: 31 - len(suffix)]
        sanitized = f"{truncated}{suffix}"
        idx += 1

    existing_names.add(sanitized)
    return sanitized


def save_excel(all_rows: list, tag_buckets: dict, path: str):
    if not all_rows:
        print("⚠️  No data to save.")
        return

    all_df = _build_df(all_rows)

    summary_rows = [
        {"Tag Label": lbl, "Question Count": len(rows)}
        for lbl, rows in tag_buckets.items()
        if rows
    ]
    summary_rows.append(
        {"Tag Label": "TOTAL (deduplicated)", "Question Count": len(all_rows)}
    )
    summary_df = pd.DataFrame(summary_rows)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        all_df.to_excel(writer, sheet_name="All Questions", index=False)
        _style_sheet(writer.sheets["All Questions"], all_df)

        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        ws_sum = writer.sheets["Summary"]
        for cell in ws_sum[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        ws_sum.column_dimensions["A"].width = 30
        ws_sum.column_dimensions["B"].width = 18

        existing_sheets = {"All Questions", "Summary"}
        for tag_label, rows in tag_buckets.items():
            if not rows:
                continue
            sheet_name = _sanitize_sheet_name(tag_label, existing_sheets)
            t_df = _build_df(rows)
            t_df.to_excel(writer, sheet_name=sheet_name, index=False)
            _style_sheet(writer.sheets[sheet_name], t_df)

    print(f"\n✅  Saved {len(all_rows)} questions --> {path}")
    non_empty = len([b for b in tag_buckets.values() if b])
    print(f"    Sheets: 'All Questions' + 'Summary' + {non_empty} tag sheets")


# ── CLI ───────────────────────────────────────────────────────────────────────


def main():
    parser = argparse.ArgumentParser(
        description="Scrape GMAT Club CR x Source:OG questions (with login)"
    )
    parser.add_argument(
        "--tags-excel",
        default=TAGS_EXCEL,
        help=f"Path to gmatclub_tags.xlsx (default: {TAGS_EXCEL})",
    )
    parser.add_argument(
        "--output",
        default=OUTPUT_EXCEL,
        help=f"Output Excel filename (default: {OUTPUT_EXCEL})",
    )
    parser.add_argument(
        "--session",
        default=SESSION_FILE,
        help=f"Session cookie file (default: {SESSION_FILE})",
    )
    parser.add_argument(
        "--fresh-login",
        action="store_true",
        help="Force re-login even if a saved session exists",
    )
    args = parser.parse_args()

    tags_excel = args.tags_excel
    if not Path(tags_excel).exists():
        alt = Path(__file__).parent / tags_excel
        if alt.exists():
            tags_excel = str(alt)
        else:
            sys.exit(
                f"Cannot find {tags_excel}. Make sure gmatclub_tags.xlsx is in the same folder."
            )

    tags = load_cr_og_tags(tags_excel)
    if not tags:
        sys.exit(
            "No CR Source:OG tags found in gmatclub_tags.xlsx.\n"
            "Check that the Category column contains 'Critical Reasoning (CR)' exactly."
        )

    try:
        import nest_asyncio

        nest_asyncio.apply()
    except ImportError:
        pass

    asyncio.run(run_scraper(tags, args.output, args.session, args.fresh_login))


if __name__ == "__main__":
    main()
