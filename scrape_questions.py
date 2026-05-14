"""
GMAT Club -- Step 3: Full Question Scraper Pipeline
=====================================================
Reads any question-links Excel file (output of Step 2 e.g. PS_OG_questions.xlsx),
visits each question page, and extracts the full question content + expert answer.

Generic Design
--------------
Just change INPUT_EXCEL at the top (or use --input flag) to point to any
Step-2 output file. The script reads the "Link" column and processes everything.

Output Columns  (matches DS_OG_questions.xlsx + gmatclub_batch_processed.csv)
------------------------------------------------------------------------------
  No            | Row number
  Title         | Question title / snippet  (from Step 2)
  Category      | e.g. Problem Solving (PS)  (from Step 2)
  Tags          | Full tag list from question page  (enriched)
  Tag IDs       | Corresponding tag IDs  (enriched)
  Difficulty    | Extracted from Tags  e.g. "555-605 (Medium)"
  Source        | All "Source: *" tags joined  e.g. "OG 2022 | OG 2025-2026"
  Link          | Direct URL to question
  Question+Text | Full cleaned question text (stem + statements)
  question-stem | Text up to the first "?"
  Statement 1   | DS statement 1  (blank for PS)
  Statement 2   | DS statement 2  (blank for PS)
  answer-by     | Expert who posted the solution
  answer-detail | Full expert solution text
  answer        | Final answer letter  e.g. "D"

Usage
-----
  python scrape_questions.py                          # uses default INPUT_EXCEL
  python scrape_questions.py --input PS_OG_questions.xlsx
  python scrape_questions.py --input DS_OG_questions.xlsx --output DS_full.xlsx
  python scrape_questions.py --input DS_OG_questions.xlsx --enrich-tags
  python scrape_questions.py --fresh-login            # force re-login
  python scrape_questions.py --start-row 51           # resume from row 51
"""

import argparse
import asyncio
import json
import re
import sys
from pathlib import Path
from urllib.parse import urlencode

try:
    from bs4 import BeautifulSoup
    import pandas as pd
    from openpyxl.styles import Font, PatternFill, Alignment
    from playwright.async_api import async_playwright
except ImportError:
    sys.exit(
        "Missing dependencies. Run:\n"
        "  pip install beautifulsoup4 pandas openpyxl playwright\n"
        "  playwright install chromium"
    )

# ── Config  (only change INPUT_EXCEL to switch question sets) ─────────────────

INPUT_EXCEL = "DS_OG_questions.xlsx"  # <-- change this for other files
OUTPUT_EXCEL = None  # auto-derived from input if not set
SESSION_FILE = "gmatclub_session.json"
BATCH_SIZE = 50  # save progress every N questions
LOGIN_URL = "https://gmatclub.com/forum/ucp.php?mode=login"
FORUM_HOME = "https://gmatclub.com/forum/"
CF_WAIT_MS = 4000
CF_RETRIES = 20
PAGE_TIMEOUT = 60000
CRAWL_DELAY = 3  # seconds between question page requests


# ── Text Cleaners (ported from original scrapper-script.py) ──────────────────


def clean_text(text: str) -> str:
    """Remove everything after 'Show Answer' button text."""
    if "Show Answer" in text:
        text = text.split("Show Answer")[0]
    return text.strip()


def extract_question_stem(text: str) -> str:
    """Return text up to and including the first '?'."""
    idx = text.find("?")
    if idx != -1:
        return text[: idx + 1].strip()
    return text.strip()


def extract_statements(text: str):
    """Extract DS Statement 1 and Statement 2 from question text."""
    s1, s2 = "", ""
    match1 = re.search(
        r"(?:1\)|Statement \(1\))(.+?)(?:2\)|Statement \(2\))", text, re.S
    )
    match2 = re.search(r"(?:2\)|Statement \(2\))(.+)", text, re.S)
    if match1:
        s1 = match1.group(1).strip()
    if match2:
        s2 = match2.group(1).strip()
    return s1, s2


def extract_answer_letter(text: str) -> str:
    """
    Try to pull the final answer letter (A-E) from expert answer detail.
    Looks for patterns like: 'Answer: D', 'The answer is D', 'OA: D'
    """
    patterns = [
        r"[Aa]nswer(?:\s+is)?[:\s]+([A-Ea-e])\b",
        r"\bOA[:\s]+([A-Ea-e])\b",
        r"[Cc]orrect\s+answer\s+is\s+([A-Ea-e])\b",
        r"[Tt]he\s+answer\s+(?:is\s+)?([A-Ea-e])\b",
    ]
    for pattern in patterns:
        m = re.search(pattern, text)
        if m:
            return m.group(1).upper()
    return ""


def extract_difficulty(tags: str) -> str:
    """Pull difficulty label from the full tags string."""
    for part in tags.split(" | "):
        p = part.strip()
        if re.search(r"\d{3}[-\u2013]\d{3}|Sub \d{3}|700\+|800\+", p):
            # Strip leading "Difficulty: " prefix if present
            return re.sub(r"^Difficulty:\s*", "", p).strip()
    return ""


def extract_sources(tags: str) -> str:
    """Pull all Source tags, strip the 'Source: ' prefix, join with ' | '."""
    sources = []
    for part in tags.split(" | "):
        p = part.strip()
        if p.startswith("Source:"):
            sources.append(p.replace("Source:", "").strip())
    return " | ".join(sources)


# ── Per-Page Extractors ───────────────────────────────────────────────────────


def extract_full_tags(html: str):
    """
    Parse div#taglist a.tag_css_link[href*='tag_id='] from question page.
    Returns (tags_str, tag_ids_str)  -- labels stripped of leading/trailing spaces.
    """
    soup = BeautifulSoup(html, "html.parser")
    taglist = soup.select("div#taglist a.tag_css_link[href*='tag_id=']")
    labels, ids = [], []
    for a in taglist:
        label = a.get_text(strip=True)  # strip leading space noted in spec
        m = re.search(r"tag_id=(\d+)", a.get("href", ""))
        if m and label:
            labels.append(label)
            ids.append(m.group(1))
    return " | ".join(labels), " | ".join(ids)


def extract_question_text(html: str):
    """Extract raw question text from first post .item.text selector."""
    soup = BeautifulSoup(html, "html.parser")
    element = soup.select_one("#posts .post-wrapper.first-post .item.text")
    if element:
        return element.get_text("\n", strip=True)
    return "Not found"


def extract_expert_answer(html: str):
    """
    Find the first post marked with .expert.box.top and return
    (expert_name, answer_detail).
    """
    soup = BeautifulSoup(html, "html.parser")
    posts = soup.select(".post-wrapper.post-separator")
    for post in posts:
        if post.select_one(".expert.box.top"):
            name_tag = post.select_one(".poster-name")
            detail_tag = post.select_one(".item.text")
            name = name_tag.get_text(strip=True) if name_tag else ""
            detail = detail_tag.get_text("\n", strip=True) if detail_tag else ""
            return name, detail
    return "no expert answer found", "no expert answer found"


# ── Cloudflare Wait ───────────────────────────────────────────────────────────


async def wait_for_cloudflare(page) -> bool:
    for attempt in range(CF_RETRIES):
        content = await page.content()
        if (
            "Just a moment" not in content
            and "security verification" not in content.lower()
        ):
            return True
        print(f"    ⏳ Cloudflare... ({attempt + 1}/{CF_RETRIES})")
        await page.wait_for_timeout(CF_WAIT_MS)
    return False


# ── Login Handler ─────────────────────────────────────────────────────────────


def is_logged_in(html: str) -> bool:
    indicators = ["ucp.php?mode=logout", "icon-svg-logout", "My Profile"]
    html_lower = html.lower()
    return any(ind.lower() in html_lower for ind in indicators)


async def login(page, session_file: str, fresh_login: bool = False) -> bool:
    session_path = Path(session_file)

    # Try restoring saved session first
    if not fresh_login and session_path.exists():
        print("\n🔑  Found saved session -- attempting to restore...")
        try:
            cookies = json.loads(session_path.read_text())
            await page.context.add_cookies(cookies)
            await page.goto(FORUM_HOME, timeout=PAGE_TIMEOUT)
            await wait_for_cloudflare(page)
            if is_logged_in(await page.content()):
                print("    ✅  Session restored -- already logged in!")
                return True
            print("    ⚠️  Saved session expired -- need to log in again.")
        except Exception as e:
            print(f"    ⚠️  Could not restore session: {e}")

    # Manual login
    print("\n🔐  Opening GMAT Club login page...")
    print("    👉  Please log in manually in the browser window.")
    print("    👉  Script continues automatically once login is detected.")
    print("    (You have up to 3 minutes)\n")
    await page.goto(LOGIN_URL, timeout=PAGE_TIMEOUT)
    await wait_for_cloudflare(page)

    for i in range(36):  # 36 x 5s = 3 minutes
        if is_logged_in(await page.content()):
            print("\n    ✅  Login detected! Starting scrape...\n")
            cookies = await page.context.cookies()
            session_path.write_text(json.dumps(cookies, indent=2))
            print(f"    💾  Session saved to {session_file}")
            return True
        remaining = (36 - i) * 5
        print(f"    ⏳  Waiting for login... ({remaining}s remaining)", end="\r")
        await page.wait_for_timeout(5000)

    print("\n    ❌  Login timeout. Please re-run the script.")
    return False


# ── Tag Enrichment ────────────────────────────────────────────────────────────


async def enrich_tags(page, rows: list) -> None:
    """
    Visit each question URL and replace Tags / Tag IDs with the
    full list from div#taglist on the question page itself.
    Mutates rows in-place.
    """
    print(f"\n🔎  Enriching tags from question pages ({len(rows)} questions)...")
    for i, row in enumerate(rows, 1):
        url = row["Link"]
        print(f"  [{i:>4}/{len(rows)}] {url}")
        try:
            await page.goto(url, timeout=PAGE_TIMEOUT)
            await page.wait_for_load_state("networkidle", timeout=15000)
            html = await page.content()
            # comment it out for login disble
            if not is_logged_in(html):
                print("    ⚠️  Session expired during tag enrichment -- stopping.")
                break

            tags_str, ids_str = extract_full_tags(html)
            if tags_str:
                row["Tags"] = tags_str
                row["Tag IDs"] = ids_str
        except Exception as e:
            print(f"    ✗ Failed: {e}")

        await asyncio.sleep(CRAWL_DELAY)


# ── Core Question Scraper ─────────────────────────────────────────────────────


async def scrape_question_page(page, url: str, row: dict) -> dict:
    """
    Visit a single question page and extract all content.
    Merges with the existing row dict from Step 2.
    """
    result = {
        # Carry over Step 2 fields
        "Title": row.get("Title", ""),
        "Category": row.get("Category", ""),
        "Tags": row.get("Tags", ""),
        "Tag IDs": row.get("Tag IDs", ""),
        "Link": url,
        # Step 3 fields (filled below)
        "Difficulty": "",
        "Source": "",
        "Question+Text": "",
        "question-stem": "",
        "Statement 1": "",
        "Statement 2": "",
        "answer-by": "",
        "answer-detail": "",
        "answer": "",
    }

    try:
        await page.goto(url, timeout=PAGE_TIMEOUT)
        await wait_for_cloudflare(page)

        try:
            await page.wait_for_selector(
                "#posts .post-wrapper.first-post .item.text", timeout=20000
            )
        except Exception:
            print(f"    ⚠️  Selector not found on: {url}")

        html = await page.content()
        # comment it out for login disble
        if not is_logged_in(html):
            print(f"    ⚠️  Session expired on: {url}")
            result["Question+Text"] = "SESSION_EXPIRED"
            return result

        # Full tags from question page (overrides search-card tags)
        tags_str, ids_str = extract_full_tags(html)
        if tags_str:
            result["Tags"] = tags_str
            result["Tag IDs"] = ids_str

        # Derived from tags
        result["Difficulty"] = extract_difficulty(result["Tags"])
        result["Source"] = extract_sources(result["Tags"])

        # Question text
        raw_text = extract_question_text(html)
        cleaned = clean_text(raw_text)
        result["Question+Text"] = cleaned
        result["question-stem"] = extract_question_stem(cleaned)
        s1, s2 = extract_statements(cleaned)
        result["Statement 1"] = s1
        result["Statement 2"] = s2

        # Expert answer
        answer_by, answer_detail = extract_expert_answer(html)
        result["answer-by"] = answer_by
        result["answer-detail"] = answer_detail
        result["answer"] = extract_answer_letter(answer_detail)

    except Exception as e:
        print(f"    ✗ Failed ({url}): {e}")
        result["Question+Text"] = f"FAILED: {e}"

    return result


# ── Batch Saver ───────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
ALT_FILL = PatternFill("solid", fgColor="EBF3FB")

COLUMN_ORDER = [
    "No",
    "Title",
    "Category",
    "Difficulty",
    "Source",
    "Tags",
    "Tag IDs",
    "Link",
    "Question+Text",
    "question-stem",
    "Statement 1",
    "Statement 2",
    "answer-by",
    "answer-detail",
    "answer",
]

COL_WIDTHS = {
    "No": 5,
    "Title": 50,
    "Category": 22,
    "Difficulty": 22,
    "Source": 35,
    "Tags": 60,
    "Tag IDs": 30,
    "Link": 30,
    "Question+Text": 60,
    "question-stem": 50,
    "Statement 1": 40,
    "Statement 2": 40,
    "answer-by": 18,
    "answer-detail": 70,
    "answer": 10,
}


def _style_sheet(ws):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    link_col = None
    for idx, cell in enumerate(ws[1], 1):
        if cell.value == "Link":
            link_col = idx

    for row_idx in range(2, ws.max_row + 1):
        if row_idx % 2 == 0:
            for cell in ws[row_idx]:
                cell.fill = ALT_FILL
        if link_col:
            cell = ws.cell(row=row_idx, column=link_col)
            if cell.value and str(cell.value).startswith("http"):
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")

    for col_cells in ws.columns:
        header = col_cells[0].value
        ws.column_dimensions[col_cells[0].column_letter].width = COL_WIDTHS.get(
            header, 20
        )

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def save_results(results: list, output_path: str) -> None:
    if not results:
        print("⚠️  No results to save.")
        return

    df = pd.DataFrame(results)
    df["No"] = range(1, len(df) + 1)

    # Ensure all expected columns exist (fill missing with empty string)
    for col in COLUMN_ORDER:
        if col not in df.columns:
            df[col] = ""

    df = df[COLUMN_ORDER]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="All Questions", index=False)
        _style_sheet(writer.sheets["All Questions"])

    print(f"  💾  {len(results)} rows --> {output_path}")


# ── Master Orchestrator ───────────────────────────────────────────────────────


async def run_pipeline(
    input_excel: str,
    output_excel: str,
    session_file: str,
    fresh_login: bool,
    enrich_tags_flag: bool,
    start_row: int,
):
    # Load input links file
    print(f"\n📂  Reading: {input_excel}")
    df_input = pd.read_excel(input_excel)

    if "Link" not in df_input.columns:
        sys.exit("❌  Input file has no 'Link' column.")

    rows = df_input.to_dict("records")

    # Apply start_row offset (1-indexed, matches "No" column)
    if start_row > 1:
        rows = rows[start_row - 1 :]
        print(f"    ▶  Resuming from row {start_row} ({len(rows)} questions remaining)")
    else:
        print(f"    ▶  {len(rows)} questions to process")

    results = []

    # Auto-resume: if output file already exists and --start-row not given, skip done rows
    output_path_obj = Path(output_excel)
    if output_path_obj.exists() and start_row == 1:
        try:
            df_existing = pd.read_excel(output_excel)
            done_links = set(df_existing["Link"].dropna().astype(str))
            rows = [r for r in rows if str(r.get("Link", "")).strip() not in done_links]
            results = df_existing.to_dict("records")
            print(
                f"    ♻  Auto-resume: {len(done_links)} already done, {len(rows)} remaining"
            )
        except Exception as e:
            print(f"    ⚠  Could not auto-resume from {output_excel}: {e}")

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=False,
            args=[
                "--no-sandbox",
                "--disable-dev-shm-usage",
                "--disable-blink-features=AutomationControlled",
            ],
        )
        context = await browser.new_context()
        page = await context.new_page()
        await page.add_init_script(
            "Object.defineProperty(navigator,'webdriver',{get:()=>undefined})"
        )
        # Login
        if not await login(page, session_file, fresh_login):
            await browser.close()
            return

        # Optional tag enrichment pass (visits each URL once just for tags)
        if enrich_tags_flag:
            await enrich_tags(page, rows)

        # Main scrape loop
        print(f"\n{'='*60}")
        print(f"  Starting question extraction ({len(rows)} questions)")
        print(f"{'='*60}\n")

        for i, row in enumerate(rows, 1):
            url = str(row.get("Link", "")).strip()
            if not url.startswith("http"):
                print(f"  [{i:>4}/{len(rows)}] Skipping invalid link: {url}")
                continue

            print(f"  [{i:>4}/{len(rows)}] {url}")

            result = await scrape_question_page(page, url, row)
            results.append(result)

            # Check for session expiry
            if result.get("Question+Text") == "SESSION_EXPIRED":
                print("\n  ⚠️  Session expired -- saving progress and stopping.")
                break

            await asyncio.sleep(CRAWL_DELAY)

            # Save checkpoint to rolling output file, then prompt to continue or stop
            if i % BATCH_SIZE == 0:
                save_results(results, output_excel)
                print(f"\n  📦  Checkpoint: {len(results)} rows saved → {output_excel}")
                answer = await asyncio.get_event_loop().run_in_executor(
                    None, input, "  ▶  Press Enter to continue, or type 'q' to stop: "
                )
                if answer.strip().lower() == "q":
                    print("  🛑  Stopped by user.")
                    break
                print()

        await browser.close()

    # Final save
    print(f"\n{'='*60}")
    print(f"  Done! {len(results)} questions extracted.")
    save_results(results, output_excel)
    print(f"{'='*60}\n")


# ── CLI ───────────────────────────────────────────────────────────────────────


def main():
    parser = argparse.ArgumentParser(
        description="Step 3: Visit each question link and extract full content.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python scrape_questions.py
  python scrape_questions.py --input PS_OG_questions.xlsx
  python scrape_questions.py --input DS_OG_questions.xlsx --output DS_full_extracted.xlsx
  python scrape_questions.py --input PS_OG_questions.xlsx --enrich-tags
  python scrape_questions.py --start-row 101   # resume after a crash at row 100
  python scrape_questions.py --fresh-login     # force re-login
        """,
    )
    parser.add_argument(
        "--input",
        default=INPUT_EXCEL,
        help=f"Input Excel with question links (default: {INPUT_EXCEL})",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="Output Excel filename (default: <input_name>_extracted.xlsx)",
    )
    parser.add_argument(
        "--session",
        default=SESSION_FILE,
        help=f"Session cookie file (default: {SESSION_FILE})",
    )
    parser.add_argument(
        "--fresh-login",
        action="store_true",
        help="Force re-login even if saved session exists",
    )
    parser.add_argument(
        "--enrich-tags",
        action="store_true",
        help="First pass: visit each question page to get full tag list from div#taglist",
    )
    parser.add_argument(
        "--start-row",
        type=int,
        default=1,
        help="Resume from this row number (1-indexed, skips earlier rows)",
    )
    args = parser.parse_args()

    # Resolve input path
    input_path = Path(args.input)
    if not input_path.exists():
        alt = Path(__file__).parent / args.input
        if alt.exists():
            input_path = alt
        else:
            sys.exit(f"❌  Input file not found: {args.input}")

    # Auto-derive output name
    output_path = args.output or input_path.stem + "_extracted.xlsx"

    print(f"\n{'='*60}")
    print(f"  GMAT Club Step 3 -- Question Extractor")
    print(f"{'='*60}")
    print(f"  Input  : {input_path}")
    print(f"  Output : {output_path}")
    print(f"  Session: {args.session}")
    print(f"  Enrich tags: {args.enrich_tags}")
    print(f"  Start row  : {args.start_row}")

    try:
        import nest_asyncio

        nest_asyncio.apply()
    except ImportError:
        pass

    asyncio.run(
        run_pipeline(
            input_excel=str(input_path),
            output_excel=output_path,
            session_file=args.session,
            fresh_login=args.fresh_login,
            enrich_tags_flag=args.enrich_tags,
            start_row=args.start_row,
        )
    )


if __name__ == "__main__":
    main()
